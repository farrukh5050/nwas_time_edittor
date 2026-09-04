"""Writes the Modified_NWAS workbook, with the status highlighting."""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


def calc_time_dif(NWAS_data, ghost_data, output_filename="Modified_NWAS_File.xlsx"):
    NWAS_data["Time"] = pd.to_datetime(
        NWAS_data["Time"], format="%H:%M:%S", errors="coerce"
    ).dt.time
    NWAS_data["Completed at Time"] = pd.to_datetime(
        NWAS_data["Completed at Time"], format="%Y-%m-%d %H:%M:%S", errors="coerce"
    ).dt.time

    def calculate_time_difference(time1, time2):
        if pd.isna(time1) or pd.isna(time2):
            return ""
        time1_seconds = time1.hour * 3600 + time1.minute * 60 + time1.second
        time2_seconds = time2.hour * 3600 + time2.minute * 60 + time2.second
        return (time2_seconds - time1_seconds) // 60

    NWAS_data["Time Difference (mins)"] = NWAS_data.apply(
        lambda row: calculate_time_difference(row["Time"], row["Completed at Time"]),
        axis=1,
    )

    def categorize_time_diff(row):
        diff = row["Time Difference (mins)"]
        category = row["Cat"]
        if diff == "":
            return "No Data"
        if diff < 0:
            if category in ["DIA", "Oncology"] and diff >= -45:
                return "Fine (Green)"
            elif category not in ["DIA", "Oncology"] and diff >= -60:
                return "Fine (Green)"
            else:
                return "Too Early (No Highlight)"
        if diff > 1:
            return "Late (Yellow)"
        return "On Time"

    NWAS_data["Status"] = NWAS_data.apply(categorize_time_diff, axis=1)

    with pd.ExcelWriter(output_filename, engine="xlsxwriter") as writer:
        NWAS_data.to_excel(writer, sheet_name="Modified_NWAS", index=False)
        ghost_data.to_excel(writer, sheet_name="Ghost_Data", index=False)

    wb = load_workbook(output_filename)
    ws = wb["Modified_NWAS"]

    green_fill = PatternFill(
        start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
    )
    yellow_fill = PatternFill(
        start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"
    )
    red_fill = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")

    status_col = ws.max_column

    for row in range(2, ws.max_row + 1):
        status_value = ws.cell(row=row, column=status_col).value
        time_value = ws[f"G{row}"].value
        completed_value = ws[f"J{row}"].value

        if status_value == "Fine (Green)":
            ws.cell(row=row, column=status_col).fill = green_fill
        elif status_value == "Late (Yellow)":
            ws.cell(row=row, column=status_col).fill = yellow_fill

        if time_value and completed_value:
            time_value = pd.to_datetime(time_value).time()
            completed_value = pd.to_datetime(completed_value).time()
            if time_value > completed_value:
                ws[f"K{row}"].fill = red_fill

    wb.save(output_filename)
    print(f"File updated successfully Choomb: '{output_filename}'")


def write_workbook(nwas_data, ghost_data, output_path, progress=None):
    """Write the workbook and report (nwas_rows, ghost_rows)."""
    if progress:
        progress("Writing workbook...")
    calc_time_dif(nwas_data, ghost_data, output_filename=output_path)
    return len(nwas_data), len(ghost_data)
